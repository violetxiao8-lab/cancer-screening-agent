# Build AgentT's multi-community ChromaDB knowledge base.

import os
import shutil
from collections import Counter

from dotenv import load_dotenv
load_dotenv()

from pptx import Presentation
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings
from langchain_community.vectorstores import Chroma
from youtube_transcript_api import (
    YouTubeTranscriptApi,
    NoTranscriptFound,
    TranscriptsDisabled,
)

# ============================================================
# CONFIGURATION
# ============================================================

DATA_DIR = "./data"
VIDEOS_DIR = "./data/videos"

# IMPORTANT:
# Build into a separate folder first.
# Do not overwrite the currently working knowledge_base yet.
VECTOR_DB_PATH = "./knowledge_base_v2"

CHINESE_PDF = os.path.join(DATA_DIR, "slides.pdf")

INDIAN_PPTX = os.path.join(
    DATA_DIR,
    "Cancer Project Indian Americans V2.pptx"
)

VIETNAMESE_PPTX = os.path.join(
    DATA_DIR,
    "Cancer Project Vietnamese Americans V2.pptx"
)

ACS_PDF = os.path.join(DATA_DIR, "aanhpi_cff.pdf")

VIDEOS = {
    "epidemiology": "j2M0tZz6LI8",
    "colorectal_screening": "Ltz4Tb1I7kQ",
    "lung_cancer_disparity": "LHjnKMx3OTY",
}

os.makedirs(VIDEOS_DIR, exist_ok=True)


# ============================================================
# HELPERS
# ============================================================

def load_pdf(path, source_name, community, source_type):
    if not os.path.exists(path):
        print(f"  ❌ Missing: {path}")
        return []

    loader = PyPDFLoader(path)
    docs = loader.load()

    for doc in docs:
        doc.metadata.update({
            "source_type": source_type,
            "source_name": source_name,
            "community": community,
        })

    print(
        f"  ✅ {source_name}: "
        f"{len(docs)} pages loaded "
        f"[community={community}]"
    )

    return docs


def extract_pptx(path, source_name, community):
    if not os.path.exists(path):
        print(f"  ❌ Missing: {path}")
        return []

    presentation = Presentation(path)
    documents = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    text_parts.append(text)

            # Extract text contained in PowerPoint tables.
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            text_parts.append(cell_text)

        slide_text = "\n".join(text_parts).strip()

        if not slide_text:
            continue

        documents.append(
            Document(
                page_content=slide_text,
                metadata={
                    "source_type": "presentation_slides",
                    "source_name": source_name,
                    "community": community,
                    "slide_number": slide_number,
                },
            )
        )

    print(
        f"  ✅ {source_name}: "
        f"{len(documents)} text-containing slides loaded "
        f"[community={community}]"
    )

    return documents


# ============================================================
# STEP 1: LOAD COMMUNITY DOCUMENTS
# ============================================================

print("\n" + "=" * 60)
print("STEP 1: Loading community documents")
print("=" * 60)

all_documents = []

# Existing Chinese-American material
all_documents.extend(
    load_pdf(
        CHINESE_PDF,
        "Cancer Screening for Chinese Americans (Slides)",
        "chinese_american",
        "presentation_slides",
    )
)

# New Indian-American material
all_documents.extend(
    extract_pptx(
        INDIAN_PPTX,
        "Cancer Screening for Indian Americans (Slides)",
        "indian_american",
    )
)

# New Vietnamese-American material
all_documents.extend(
    extract_pptx(
        VIETNAMESE_PPTX,
        "Cancer Screening for Vietnamese Americans (Slides)",
        "vietnamese_american",
    )
)

# ACS document is useful across populations.
all_documents.extend(
    load_pdf(
        ACS_PDF,
        "ACS AANHPI Cancer Facts and Figures",
        "general",
        "research_document",
    )
)


# ============================================================
# STEP 2: LOAD/FETCH EXISTING VIDEO TRANSCRIPTS
# ============================================================

print("\n" + "=" * 60)
print("STEP 2: Loading YouTube transcripts")
print("=" * 60)

ytt_api = YouTubeTranscriptApi()

for topic, video_id in VIDEOS.items():
    saved_path = os.path.join(
        VIDEOS_DIR,
        f"{topic}.txt"
    )

    transcript_text = None

    if os.path.exists(saved_path):
        print(f"  📂 {topic}: Loading saved transcript...")

        with open(
            saved_path,
            "r",
            encoding="utf-8"
        ) as f:
            transcript_text = f.read().strip()

        print(
            f"      ✅ Loaded "
            f"({len(transcript_text):,} chars)"
        )

    else:
        print(
            f"  🌐 {topic} ({video_id}): "
            f"Fetching from YouTube..."
        )

        try:
            fetched = ytt_api.fetch(
                video_id,
                languages=("en", "en-US")
            )

            snippets = [
                s.text.strip()
                for s in fetched
                if s.text.strip()
            ]

            transcript_text = " ".join(snippets)

            if not transcript_text:
                raise ValueError("Empty transcript")

            with open(
                saved_path,
                "w",
                encoding="utf-8"
            ) as f:
                f.write(transcript_text)

            print(
                f"      ✅ Fetched and saved "
                f"({len(transcript_text):,} chars)"
            )

        except (
            NoTranscriptFound,
            TranscriptsDisabled
        ) as exc:
            print(
                f"      ❌ No transcript available: {exc}"
            )
            transcript_text = None

        except Exception as exc:
            print(f"      ❌ Failed: {exc}")
            transcript_text = None

    if transcript_text:
        all_documents.append(
            Document(
                page_content=transcript_text,
                metadata={
                    "source_type": "video_transcript",
                    "source_name": (
                        "YouTube: "
                        + topic.replace("_", " ").title()
                    ),
                    "community": "general",
                    "video_id": video_id,
                    "topic": topic,
                },
            )
        )


# ============================================================
# STEP 3: VERIFY DOCUMENT DISTRIBUTION
# ============================================================

print("\n" + "=" * 60)
print("STEP 3: Checking document distribution")
print("=" * 60)

community_document_counts = Counter(
    doc.metadata.get("community", "unknown")
    for doc in all_documents
)

for community, count in sorted(
    community_document_counts.items()
):
    print(
        f"  {community}: "
        f"{count} documents/slides"
    )

required_communities = {
    "chinese_american",
    "indian_american",
    "vietnamese_american",
}

missing_communities = (
    required_communities
    - set(community_document_counts.keys())
)

if missing_communities:
    raise RuntimeError(
        "Missing required community content: "
        + ", ".join(sorted(missing_communities))
    )


# ============================================================
# STEP 4: SPLIT DOCUMENTS
# ============================================================

print("\n" + "=" * 60)
print("STEP 4: Splitting documents into chunks")
print("=" * 60)

splitter = RecursiveCharacterTextSplitter(
    chunk_size=1000,
    chunk_overlap=150,
    separators=[
        "\n\n",
        "\n",
        ". ",
        "! ",
        "? ",
        " ",
        "",
    ],
)

chunks = splitter.split_documents(all_documents)

print(
    f"  ✅ Created {len(chunks)} chunks "
    f"from {len(all_documents)} documents"
)

community_chunk_counts = Counter(
    chunk.metadata.get("community", "unknown")
    for chunk in chunks
)

print("\n  Chunks by community:")

for community, count in sorted(
    community_chunk_counts.items()
):
    print(f"    {community}: {count}")

source_counts = Counter(
    chunk.metadata.get("source_type", "unknown")
    for chunk in chunks
)

print("\n  Chunks by source type:")

for source, count in sorted(source_counts.items()):
    print(f"    {source}: {count}")


# ============================================================
# STEP 5: BUILD A CLEAN TEST DATABASE
# ============================================================

print("\n" + "=" * 60)
print("STEP 5: Building clean ChromaDB")
print("=" * 60)

api_key = os.environ.get("OPENAI_API_KEY")

if not api_key:
    raise RuntimeError(
        "OPENAI_API_KEY was not found. "
        "Add it to your .env file before running this script."
    )

# We are rebuilding knowledge_base_v2 from scratch.
# Existing production knowledge_base is NOT touched.
if os.path.exists(VECTOR_DB_PATH):
    print(
        f"  🧹 Removing previous test database: "
        f"{VECTOR_DB_PATH}"
    )
    shutil.rmtree(VECTOR_DB_PATH)

embeddings = OpenAIEmbeddings(
    model="text-embedding-3-small",
    openai_api_key=api_key,
)

print(
    "  ⏳ Embedding documents using "
    "text-embedding-3-small..."
)

vectorstore = Chroma.from_documents(
    documents=chunks,
    embedding=embeddings,
    persist_directory=VECTOR_DB_PATH,
)

print(
    f"  ✅ New ChromaDB saved to: "
    f"{VECTOR_DB_PATH}"
)


# ============================================================
# STEP 6: COMMUNITY-SPECIFIC RETRIEVAL TESTS
# ============================================================

print("\n" + "=" * 60)
print("STEP 6: Testing community-specific retrieval")
print("=" * 60)

tests = [
    (
        "chinese_american",
        "What is the most common cancer "
        "among Chinese American women?",
    ),
    (
        "indian_american",
        "What is the most common cancer "
        "among Indian American men?",
    ),
    (
        "vietnamese_american",
        "What is the most common cancer "
        "among Vietnamese American men?",
    ),
]

for community, query in tests:

    print("\n" + "-" * 60)
    print(f"Community: {community}")
    print(f"Question: {query}")

    community_results = vectorstore.similarity_search(
        query,
        k=3,
        filter={"community": community},
    )

    general_results = vectorstore.similarity_search(
        query,
        k=2,
        filter={"community": "general"},
    )

    results = community_results + general_results

    if not results:
        print("  ❌ No retrieval results")
        continue

    for index, result in enumerate(
        results,
        start=1
    ):
        print(
            f"\n  Result {index}"
        )

        print(
            "  Source:",
            result.metadata.get(
                "source_name",
                "unknown"
            )
        )

        print(
            "  Community:",
            result.metadata.get(
                "community",
                "unknown"
            )
        )

        print(
            "  Preview:",
            result.page_content[:250]
            .replace("\n", " "),
            "..."
        )


# ============================================================
# COMPLETE
# ============================================================

print("\n" + "=" * 60)
print("✅ MULTI-COMMUNITY KNOWLEDGE BASE READY")
print("=" * 60)

print(
    "\nBuilt communities:"
    "\n  • Chinese American"
    "\n  • Indian American"
    "\n  • Vietnamese American"
    "\n  • General AANHPI/supporting sources"
)

print(
    f"\nTest database location: {VECTOR_DB_PATH}"
)

print(
    "\nIMPORTANT: Your existing ./knowledge_base "
    "has NOT been replaced."
)